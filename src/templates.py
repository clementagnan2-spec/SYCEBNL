# -*- coding: utf-8 -*-
"""
templates.py
============
Crée les fichiers modèles Excel que l'utilisateur remplit puis importe dans
le logiciel : modèle de balance des comptes, et modèle de budget (mode Projet).
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"
HEADER_FONT = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
LEGEND_FONT = Font(name=FONT_NAME, size=9, italic=True, color="808080")
EXAMPLE_FILL = PatternFill("solid", fgColor="FFF2CC")


def create_balance_template(path, period_label="N"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance"
    period_label = str(period_label).strip().upper() or "N"
    ws["A1"] = f"Modèle de balance des comptes SYCEBNL — Balance {period_label} — à importer dans le logiciel"
    ws["A1"].font = Font(name=FONT_NAME, size=12, bold=True)
    ws["A2"] = ("Renseignez une ligne par compte utilisé (ne pas insérer de ligne de "
                "total). Les colonnes peuvent être réordonnées mais leurs en-têtes "
                "doivent rester reconnaissables.")
    ws["A2"].font = LEGEND_FONT
    ws.merge_cells("A2:E2")
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = ["Compte", "Intitulé", "Débit", "Crédit"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    example = ["521", "Banque XYZ - Compte courant", 1500000, 0]
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=5, column=c, value=v)
        cell.fill = EXAMPLE_FILL
        cell.font = Font(name=FONT_NAME, italic=True, size=10)
    ws.cell(row=5, column=5, value="← exemple, à remplacer/supprimer").font = LEGEND_FONT

    for c, w in enumerate([12, 45, 16, 16, 32], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    return path


def create_budget_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws["A1"] = "Modèle de budget — mode Projet de développement"
    ws["A1"].font = Font(name=FONT_NAME, size=12, bold=True)
    ws["A2"] = ("Une ligne par rubrique budgétaire. La colonne 'Comptes' liste les "
                "préfixes de comptes SYCEBNL rattachés à cette rubrique, séparés par "
                "un point-virgule (ex : 60;61;62).")
    ws["A2"].font = LEGEND_FONT
    ws.merge_cells("A2:D2")
    ws["A2"].alignment = Alignment(wrap_text=True)

    headers = ["Rubrique", "Comptes", "Budget"]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=c, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    example = ["Formation des bénéficiaires", "618;661", 5000000]
    for c, v in enumerate(example, start=1):
        cell = ws.cell(row=5, column=c, value=v)
        cell.fill = EXAMPLE_FILL
        cell.font = Font(name=FONT_NAME, italic=True, size=10)
    ws.cell(row=5, column=4, value="← exemple, à remplacer/supprimer").font = LEGEND_FONT

    for c, w in enumerate([35, 25, 16, 32], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A5"
    wb.save(path)
    return path
